"""Download DOTAX 'Tax Credits Claimed' claim COUNTS by income class, TY2018-2023.

Outputs packages/tax_modeler/src/tax_modeler/data/raw/dotax_credit_claims_by_agi.csv,
consumed by tax_modeler.scenarios.quintile_analysis to impute which households
claim the renewable energy technologies credit (REEC) and the capital goods
excise tax credit (CGEC), so a credit cut lands on claimants rather than being
spread over every filer in an income class.

Companion to scripts/fetch_dotax_credits_historical.py, which pulls the REEC
dollar amounts only. URL patterns (verified September 24, 2026):
    TY2018-2022: https://files.hawaii.gov/tax/stats/stats/credits/archive/<YEAR>credit_tables.xlsx
    TY2023:      https://files.hawaii.gov/tax/stats/stats/credits/<YEAR>credit_tables.xlsx
(the newest year sits outside archive/ until the next edition is published).

Sheets parsed:
    Table 2 — number of individual income tax returns by Hawaii AGI class
              (Forms N-11 + N-15; the claim-rate denominator)
    A-1     — dollar amounts by taxpayer type (individual vs all, per credit)
    A-5     — dollar amounts claimed by individuals, by income class ($1,000)
    A-6     — number of credits claimed on individual returns, by income class

One row per (year, credit, agi_bin). "d" (suppressed) cells are recorded as
blank with suppressed=True.
"""
from __future__ import annotations

import csv
import sys
import urllib.error
import urllib.request
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
OUT_CSV = (REPO_ROOT / "packages" / "tax_modeler" / "src" / "tax_modeler" / "data" / "raw"
           / "dotax_credit_claims_by_agi.csv")
DL_DIR = Path("/tmp/dotax_credits")

URLS = (
    "https://files.hawaii.gov/tax/stats/stats/credits/archive/{year}credit_tables.xlsx",
    "https://files.hawaii.gov/tax/stats/stats/credits/{year}credit_tables.xlsx",
)
YEARS = [2018, 2019, 2020, 2021, 2022, 2023]
AGI_BINS = ["<$10K", "$10K-$30K", "$30K-$60K", "$60K-$100K", "$100K-$200K", "$200K+"]
CREDITS = {  # key -> row-title substring (titles drift slightly across editions)
    "reec": "enewable Energy Technologies",
    "cgec": "apital Goods Excise",
}


def _download(year: int) -> Path:
    DL_DIR.mkdir(parents=True, exist_ok=True)
    dest = DL_DIR / f"{year}.xlsx"
    if dest.exists() and dest.stat().st_size > 10_000:
        return dest
    for fmt in URLS:
        try:
            urllib.request.urlretrieve(fmt.format(year=year), dest)
            return dest
        except urllib.error.HTTPError:
            continue
    raise FileNotFoundError(f"TY{year} credit tables not found at any known URL")


def _num(cell) -> float | None:
    s = "" if cell is None else str(cell).strip()
    try:
        return float(s)
    except ValueError:
        return None       # 'd', '-', 'na', blank


def _row(ws, needle: str) -> tuple:
    for r in ws.iter_rows(values_only=True):
        if r and r[0] and needle in str(r[0]):
            return r
    raise LookupError(f"no row containing {needle!r} in {ws.title}")


def parse_year(year: int) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(_download(year), read_only=True, data_only=True)
    returns = []
    for label in ("Less than $10,000", "$10,000 to $29,999", "$30,000 to $59,999",
                  "$60,000 to $99,999", "$100,000 to $199,999", "$200,000 or more"):
        # The label's column moved in TY2022; "All Individual Returns" is the
        # last number on the row in every edition.
        row = next(r for r in wb["Table 2"].iter_rows(values_only=True)
                   if any(c is not None and str(c).strip() == label for c in r))
        returns.append([v for v in map(_num, row) if v is not None][-1])

    rows = []
    for credit, needle in CREDITS.items():
        a1 = _row(wb["A-1"], needle)
        all_k, ind_k = _num(a1[1]), _num(a1[2])
        amt = _row(wb["A-5"], needle)
        cnt = _row(wb["A-6"], needle)
        for i, agi_bin in enumerate(AGI_BINS):
            a, c = _num(amt[2 + i]), _num(cnt[2 + i])
            rows.append({
                "year": year, "credit": credit, "agi_bin": agi_bin,
                "claims": c, "amount_$K": a, "returns": returns[i],
                "credit_all_taxpayers_$K": all_k, "credit_individuals_$K": ind_k,
                "suppressed": a is None or c is None,
            })
    return rows


def main() -> int:
    rows = []
    for y in YEARS:
        try:
            got = parse_year(y)
        except Exception as e:  # noqa: BLE001 — report and keep going
            print(f"TY{y} FAILED: {e}", file=sys.stderr)
            continue
        rows += got
        reec = [r for r in got if r["credit"] == "reec"]
        print(f"TY{y}: REEC claims {sum(r['claims'] or 0 for r in reec):,.0f} "
              f"on {sum(r['returns'] or 0 for r in reec):,.0f} returns")
    if not rows:
        return 1
    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUT_CSV.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    print(f"Wrote {OUT_CSV.relative_to(REPO_ROOT)} ({len(rows)} rows)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
