"""Parse IRS SOI Table 2 for Hawaiʻi: returns, wages and AGI by size of adjusted gross income.

The IRS publishes one workbook per tax year, "Table 2. Individual Income and Tax Data, by
State and Size of Adjusted Gross Income", for each state at

    https://www.irs.gov/pub/irs-soi/{YY}in12hi.xlsx        (YY = two-digit tax year)

It counts every federal return filed from a Hawaiʻi address, in ten AGI classes, so it includes
military families and people who moved in or out during the year. The Department of Taxation's
resident-return tables do not, which is why the two benchmarks differ (see
``tax_modeler.agreement``). Money amounts are in thousands of dollars.

Usage:
    .venv/bin/python scripts/parse_irs_soi_hi_table2.py --download 2022 2023
    .venv/bin/python scripts/parse_irs_soi_hi_table2.py 2023=PATH/23in12hi.xlsx 2022=PATH/22in12hi.xlsx

Writes packages/tax_modeler/src/tax_modeler/data/calibration/irs_soi_hi_table2.json with each
file's SHA-256. Each year is checked before it is written: the ten classes must add to the
printed all-returns column (returns are rounded to the nearest 10, so within 50), and the
wage counts and amounts and AGI must add to theirs.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
import tempfile
import urllib.request
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
OUT = REPO / "packages" / "tax_modeler" / "src" / "tax_modeler" / "data" / "calibration" / "irs_soi_hi_table2.json"
URL = "https://www.irs.gov/pub/irs-soi/{yy}in12hi.xlsx"

#: The ten classes, left edge inclusive, in the order the workbook prints them.
CLASSES = [
    ("Under $1", None, 1), ("$1 to $10,000", 1, 10_000), ("$10,000 to $25,000", 10_000, 25_000),
    ("$25,000 to $50,000", 25_000, 50_000), ("$50,000 to $75,000", 50_000, 75_000),
    ("$75,000 to $100,000", 75_000, 100_000), ("$100,000 to $200,000", 100_000, 200_000),
    ("$200,000 to $500,000", 200_000, 500_000), ("$500,000 to $1,000,000", 500_000, 1_000_000),
    ("$1,000,000 or more", 1_000_000, None),
]


def _label(cell) -> str:
    return " ".join(str(cell or "").split())


def parse_workbook(path: Path) -> dict:
    """The rows this model uses from one year's workbook, checked against their printed totals."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    title = _label(rows[0][0])
    if "Table 2" not in title or "State" not in title:
        raise ValueError(f"{path.name}: not IRS Table 2 by state ({title!r})")
    # The 2023 workbook labels the state in its header; the 2022 one does not, and is identified by
    # the IRS's own file name (in12hi = Hawaiʻi). One of the two must say Hawaii.
    labelled = any(_label(c) == "HAWAII" for r in rows[:8] for c in r[:3])
    if not labelled and not path.name.lower().endswith("in12hi.xlsx"):
        raise ValueError(f"{path.name}: nothing says this is the Hawaii table")
    year = int(title.rsplit(" ", 1)[-1])

    def row_of(prefix: str, after: int = 0) -> int:
        for i in range(after, len(rows)):
            if _label(rows[i][0]).startswith(prefix):
                return i
        raise ValueError(f"{path.name}: no row starting {prefix!r}")

    def values(i: int) -> tuple[float, list[float]]:
        cells = list(rows[i][1:12])
        if len(cells) < 11 or any(c is None for c in cells):
            raise ValueError(f"{path.name}: row {i} ({_label(rows[i][0])!r}) is incomplete")
        return float(cells[0]), [float(c) for c in cells[1:]]

    r_ret = row_of("Number of returns")
    r_joint = row_of("Number of joint returns")
    r_agi = row_of("Adjusted gross income (AGI)")
    r_wn = row_of("Salaries and wages in AGI: Number")
    if _label(rows[r_wn + 1][0]) != "Amount":
        raise ValueError(f"{path.name}: the row under the wage count is not the amount")
    out = {"tax_year": year}
    for key, i, tol_of in (("returns", r_ret, lambda t: 50), ("joint_returns", r_joint, lambda t: 50),
                           ("agi_amount_k", r_agi, lambda t: max(10, 1e-4 * abs(t))),
                           ("wage_returns", r_wn, lambda t: 50),
                           ("wage_amount_k", r_wn + 1, lambda t: max(10, 1e-4 * abs(t)))):
        total, classes = values(i)
        if abs(sum(classes) - total) > tol_of(total):
            raise ValueError(f"{path.name}: {key} classes add to {sum(classes):,.0f}, printed total {total:,.0f}")
        out[key] = classes
        out[f"{key}_total"] = total
    return out


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("editions", nargs="*", help="YEAR=PATH to a downloaded workbook")
    ap.add_argument("--download", nargs="+", type=int, metavar="YEAR", help="fetch these tax years from irs.gov")
    ap.add_argument("--out", type=Path, default=OUT)
    args = ap.parse_args(argv)

    files: dict[int, Path] = {}
    tmp = Path(tempfile.mkdtemp(prefix="irs_table2_"))
    for year in args.download or []:
        dest = tmp / f"{year % 100:02d}in12hi.xlsx"
        urllib.request.urlretrieve(URL.format(yy=f"{year % 100:02d}"), dest)  # noqa: S310 - fixed irs.gov URL
        files[year] = dest
    for spec in args.editions:
        year, _, path = spec.partition("=")
        files[int(year)] = Path(path).expanduser()
    if not files:
        ap.error("give YEAR=PATH workbooks or --download YEAR ...")

    years: dict[str, dict] = {}
    source_files: dict[str, dict] = {}
    for year in sorted(files):
        parsed = parse_workbook(files[year])
        if parsed["tax_year"] != year:
            raise SystemExit(f"{files[year].name} is tax year {parsed['tax_year']}, not {year}")
        years[str(year)] = parsed
        source_files[str(year)] = {"file": files[year].name, "url": URL.format(yy=f"{year % 100:02d}"),
                                   "sha256": sha256(files[year])}
    doc = {
        "source": {
            "title": "Table 2. Individual Income and Tax Data, by State and Size of Adjusted Gross Income (Hawaii)",
            "publisher": "Internal Revenue Service, Statistics of Income",
            "page": "https://www.irs.gov/statistics/soi-tax-stats-individual-income-tax-statistics-historic-table-2",
            "files": source_files,
            "note": "Every federal return filed from a Hawaii address, by AGI class; money amounts in "
                    "thousands of dollars; returns rounded to the nearest 10. Includes military families and "
                    "part-year residents, unlike the Department of Taxation's resident-return tables.",
        },
        "classes": [{"label": lab, "lo": lo, "hi": hi} for lab, lo, hi in CLASSES],
        "years": years,
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(doc, indent=1, ensure_ascii=False) + "\n")
    for y, d in years.items():
        print(f"TY{y}: {d['returns_total']:,.0f} returns, wages ${d['wage_amount_k_total'] / 1e6:,.2f}B, AGI ${d['agi_amount_k_total'] / 1e6:,.2f}B")
    print(f"wrote {args.out.relative_to(REPO)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
