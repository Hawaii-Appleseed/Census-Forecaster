"""Tests for scripts/refresh_annual_anchors.py — parse, guard, failure posture.

No network: every fetcher is monkeypatched.
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from datetime import date

import openpyxl
import pytest

from census_forecaster.acs.sources.base import AnchorSource
from census_forecaster.scripts import refresh_annual_anchors as ra
from census_forecaster.scripts import refresh_national_macro as nm

SPECS = {s["filename"]: s for s in ra.ANCHOR_SPECS}


# ---------------------------------------------------------------------------
# Value builders
# ---------------------------------------------------------------------------

def test_fred_annual_keeps_years_from_first_year_and_rounds():
    rows = [{"date": "2009-01-01", "value": 80.0},
            {"date": "2010-01-01", "value": 87.4211},
            {"date": "2011-01-01", "value": 89.7}]
    assert ra.annual_from_fred_annual(rows, first_year=2010) == {2010: 87.421, 2011: 89.7}


def test_fred_quarterly_is_the_mean_of_four_quarters_complete_years_only():
    rows = [{"date": f"2023-{m:02d}-01", "value": v}
            for m, v in ((1, 100.0), (4, 110.0), (7, 120.0), (10, 130.0))]
    rows += [{"date": "2024-01-01", "value": 140.0},   # 2024 has one quarter: dropped
             {"date": "2022-01-01", "value": 1.0}]      # 2022 has one quarter: dropped
    assert ra.annual_from_fred_quarterly(rows, first_year=2010) == {2023: 115.0}


def test_fred_quarterly_with_unexpected_months_raises():
    rows = [{"date": f"2023-{m:02d}-01", "value": 1.0} for m in (2, 5, 8, 11)]
    with pytest.raises(ValueError, match="layout change"):
        ra.annual_from_fred_quarterly(rows)


def test_bls_annual_takes_only_a01_periods():
    pts = [{"year": 2024, "period": "A01", "value": 67366.0},
           {"year": 2024, "period": "Q04", "value": 63090.0},
           {"year": 2024, "period": "M03", "value": 632283.0},
           {"year": 2009, "period": "A01", "value": 40000.0}]
    assert ra.annual_from_bls(pts, first_year=2010) == {2024: 67366.0}


def test_implausible_values_raise():
    # Employment levels (~600,000) where pay (~65,000) was expected: the wrong
    # QCEW data type, which is exactly what the file's old series id pointed at.
    with pytest.raises(ValueError, match="outside"):
        ra.check_plausible("qcew_hawaii_wages.json", {2024: 632283.0})
    ra.check_plausible("qcew_hawaii_wages.json", {2024: 67366.0})


def test_years_lost_names_committed_years_the_fetch_lacks():
    assert ra.years_lost({2022: 1.0, 2023: 1.0, 2024: 1.0}, {2022: 2.0, 2023: 2.0}) == [2024]
    assert ra.years_lost({2022: 1.0}, {2022: 9.0, 2025: 9.0}) == []


# ---------------------------------------------------------------------------
# Payload schema (what AnchorSource.from_file reads)
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("name", sorted(SPECS))
def test_payload_loads_through_anchor_source(tmp_path, name):
    spec = SPECS[name]
    lo, hi = ra._RANGES[name]
    values = {y: lo + (hi - lo) * 0.1 * (1.03 ** (y - 2010)) for y in range(2010, 2026)}
    payload = ra.build_payload(spec, values, today=date(2026, 10, 7))
    assert payload["last_refresh"] == "2026-10"
    assert list(payload["values_by_year"]) == [str(y) for y in range(2010, 2026)]
    path = tmp_path / name
    path.write_text(json.dumps(payload))
    src = AnchorSource.from_file(path, publication_lag_years=1,
                                 indicator_affinity=("B19013_001E",))
    assert src.metadata["source"] == spec["source"]


# ---------------------------------------------------------------------------
# BLS chunking
# ---------------------------------------------------------------------------

def test_bls_fetch_chunks_by_window_and_concatenates(monkeypatch):
    calls = []

    def fake(series_ids, start_year=None, end_year=None, api_key=None, **kw):
        calls.append((start_year, end_year))
        return {series_ids[0]: [{"year": start_year, "period": "A01", "value": 1.0}]}

    monkeypatch.setattr(ra, "fetch_cpi_data", fake)
    pts = ra.fetch_bls_annual("S", None, first_year=2010, last_year=2026)
    assert calls == [(2010, 2019), (2020, 2026)] and len(pts) == 2     # keyless: 10-year windows
    calls.clear()
    ra.fetch_bls_annual("S", "KEY", first_year=2010, last_year=2026)
    assert calls == [(2010, 2026)]                                      # keyed: 20-year windows


# ---------------------------------------------------------------------------
# main(): writes, and the failure posture
# ---------------------------------------------------------------------------

def _good(spec_values):
    def fake_fetch_values(spec, api_key, first_year=ra.FIRST_YEAR):
        return dict(spec_values[spec["filename"]])
    return fake_fetch_values


def _values_for(name, last_year=2025):
    lo, hi = ra._RANGES[name]
    return {y: lo + (hi - lo) * 0.1 for y in range(2010, last_year + 1)}


def _old_file(tmp_path, name, last_year=2024, last_refresh="2026-01", value=1.0):
    spec = SPECS[name]
    p = ra.build_payload(spec, {y: value for y in range(2010, last_year + 1)},
                         today=date(2026, 1, 15))
    p["last_refresh"] = last_refresh
    (tmp_path / name).write_text(json.dumps(p))


def test_main_writes_every_anchor_with_todays_last_refresh(tmp_path, monkeypatch):
    for name in SPECS:
        _old_file(tmp_path, name)
    monkeypatch.setattr(ra, "fetch_values", _good({n: _values_for(n) for n in SPECS}))
    monkeypatch.delenv("BLS_API_KEY", raising=False)
    assert ra.main(["--out", str(tmp_path)]) == 0
    for name in SPECS:
        d = json.loads((tmp_path / name).read_text())
        assert d["last_refresh"] == date.today().strftime("%Y-%m")
        assert sorted(d["values_by_year"]) == [str(y) for y in range(2010, 2026)]   # 2025 added


def test_a_failed_fetch_leaves_the_file_and_its_last_refresh(tmp_path, monkeypatch, capsys):
    for name in SPECS:
        _old_file(tmp_path, name)
    before = (tmp_path / "qcew_hawaii_wages.json").read_text()
    good = _good({n: _values_for(n) for n in SPECS})

    def fake(spec, api_key, first_year=ra.FIRST_YEAR):
        if spec["filename"] == "qcew_hawaii_wages.json":
            raise RuntimeError("BLS API error: REQUEST_NOT_PROCESSED")
        return good(spec, api_key, first_year)

    monkeypatch.setattr(ra, "fetch_values", fake)
    assert ra.main(["--out", str(tmp_path)]) == 0                       # degrade, don't fail CI
    assert (tmp_path / "qcew_hawaii_wages.json").read_text() == before  # untouched, still "2026-01"
    assert json.loads((tmp_path / "pce_deflator.json").read_text())["last_refresh"] != "2026-01"
    err = capsys.readouterr().err
    assert "::warning::qcew_hawaii_wages.json: fetch failed" in err
    assert f"1 of {len(SPECS)} anchors not refreshed: qcew_hawaii_wages.json" in err
    # --strict turns the same failure into a non-zero exit
    assert ra.main(["--out", str(tmp_path), "--strict"]) == 1


def test_a_fetch_that_lacks_committed_years_is_refused(tmp_path, monkeypatch, capsys):
    name = "fred_hi_hpi.json"
    _old_file(tmp_path, name, last_year=2024)
    before = (tmp_path / name).read_text()
    short = {y: _values_for(name)[y] for y in range(2010, 2022)}        # lacks 2022-2024
    monkeypatch.setattr(ra, "ANCHOR_SPECS", (SPECS[name],))
    monkeypatch.setattr(ra, "fetch_values", lambda spec, key, first_year=ra.FIRST_YEAR: short)
    assert ra.main(["--out", str(tmp_path)]) == 0
    assert (tmp_path / name).read_text() == before
    assert "lacks 3 committed year(s) (2022-2024)" in capsys.readouterr().err


def test_dry_run_does_no_network_and_writes_nothing(tmp_path, monkeypatch):
    def boom(*a, **k):
        raise AssertionError("network during --dry-run")
    monkeypatch.setattr(ra, "fetch_values", boom)
    assert ra.main(["--out", str(tmp_path), "--dry-run"]) == 0
    assert list(tmp_path.iterdir()) == []


def test_fred_fetch_is_the_shared_ua_free_fetcher():
    """FRED stalls custom User-Agents from datacenter IPs (see
    refresh_national_macro.fetch_fred_csv). Reusing that function is what keeps
    these anchors off the same trap; a private copy could reintroduce it."""
    assert ra.fetch_fred_csv is nm.fetch_fred_csv


# ---------------------------------------------------------------------------
# HUD FMR workbook
# ---------------------------------------------------------------------------

HUD_ROWS = {2009: 1631, 2010: 1704, 2011: 1702, 2012: 1767, 2013: 1833, 2014: 1820,
            2015: 1810, 2016: 1985, 2017: 1982, 2018: 2031}


def _area(year: int) -> str:
    return "METRO26180M26180" if year <= 2015 else "METRO46520M46520"


def _hud_workbook(honolulu, *, area=_area, with_fips=True, dup=False, core_date=None) -> bytes:
    """A miniature of HUD's ``FMR_2Bed_1983_<year>.xlsx``: ``fips`` and descriptive
    columns, then ``msaYY, fmrYY_2, fmrYY`` per fiscal year (newest first), one row
    per FMR area. Hawaii County rides along as a decoy with different numbers, and
    ``fmrYY`` carries the percentile (40, or 50 for FY2012-FY2017) as in the real file."""
    years = sorted(honolulu, reverse=True)
    header = ["fips" if with_fips else "id", "census_region", "state", "county", "cousub", "name"]
    for y in years:
        header += [f"msa{y % 100:02d}", f"fmr{y % 100:02d}_2", f"fmr{y % 100:02d}"]

    def row(fips, name, values, code):
        r = [fips, 4, "15", fips[2:5], "99999", name]
        for y in years:
            r += [code(y), values[y], 50 if 2012 <= y <= 2017 else 40]
        return r

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    ws.append(row("1500199999", "Hawaii County", {y: 1111 for y in years}, lambda y: "NCNTY15001N15001"))
    ws.append(row("1500399999", "Honolulu County", honolulu, area))
    if dup:
        ws.append(row("1500399999", "Honolulu County", honolulu, area))
    buf = io.BytesIO()
    wb.save(buf)
    return _with_core_date(buf.getvalue(), core_date) if core_date else buf.getvalue()


def _with_core_date(data: bytes, value: str) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            blob = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                blob = re.sub(rb"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)",
                              rb"\g<1>" + value.encode() + rb"\g<2>", blob)
            zout.writestr(item.filename, blob)
    return out.getvalue()


def test_hud_workbook_takes_the_honolulu_row_and_the_2br_column_from_first_year():
    got = ra.annual_from_hud_workbook(_hud_workbook(HUD_ROWS))
    assert got == {y: v for y, v in HUD_ROWS.items() if y >= 2010}      # not Hawaii County's 1111
    assert all(isinstance(v, int) for v in got.values())                  # dollars, not the 40/50 percentile


def test_hud_workbook_accepts_both_honolulu_area_codes_and_rejects_a_new_one():
    ra.annual_from_hud_workbook(_hud_workbook(HUD_ROWS))                  # 26180 through FY2015, 46520 after
    recoded = lambda y: "METRO99999M99999" if y == 2018 else _area(y)     # noqa: E731
    with pytest.raises(ValueError, match="FMR area code 'METRO99999M99999'"):
        ra.annual_from_hud_workbook(_hud_workbook(HUD_ROWS, area=recoded))


def test_hud_workbook_layout_changes_raise():
    with pytest.raises(ValueError, match="no 'fips' column"):
        ra.annual_from_hud_workbook(_hud_workbook(HUD_ROWS, with_fips=False))
    with pytest.raises(ValueError, match="0 rows for fips 1500399998"):
        ra.annual_from_hud_workbook(_hud_workbook(HUD_ROWS), fips="1500399998")
    with pytest.raises(ValueError, match="2 rows for fips 1500399999"):
        ra.annual_from_hud_workbook(_hud_workbook(HUD_ROWS, dup=True))


def test_hud_workbook_fractional_dollars_and_year_gaps_raise():
    with pytest.raises(ValueError, match="not whole dollars"):
        ra.annual_from_hud_workbook(_hud_workbook({**HUD_ROWS, 2018: 2031.5}))
    with pytest.raises(ValueError, match="not one contiguous run"):
        ra.annual_from_hud_workbook(_hud_workbook({**HUD_ROWS, 2014: None}))


@pytest.mark.parametrize("bad", ["not-a-date", "2026-08-31"])
def test_hud_workbook_with_a_malformed_core_date_still_parses(bad):
    """HUD has shipped workbooks openpyxl cannot open for their dcterms dates
    (ValueError for garbage, TypeError for a bare date). The dates are metadata."""
    data = _hud_workbook(HUD_ROWS, core_date=bad)
    stripped = _hud_workbook(HUD_ROWS, core_date="2026-08-31T00:00:00Z")
    assert ra.annual_from_hud_workbook(data) == ra.annual_from_hud_workbook(stripped)
    wb = openpyxl.load_workbook(io.BytesIO(ra._strip_core_dates(data)), read_only=True)
    assert wb.sheetnames == ["Sheet"]                                     # opens once the dates are gone
    wb.close()


class _Resp:
    def __init__(self, status_code, content=b""):
        self.status_code, self.content = status_code, content


def test_hud_fetch_tries_next_year_first_and_sends_a_browser_user_agent(monkeypatch):
    calls = []
    book = _Resp(200, b"PK\x03\x04 a workbook")

    def get(url, headers=None, timeout=None):
        calls.append((url, headers))
        return _Resp(404, b"<html>not published</html>") if "2027" in url else book

    monkeypatch.setattr(ra.requests, "get", get)
    assert ra.fetch_hud_workbook(today=date(2026, 1, 5)) == book.content
    assert [u.rsplit("_", 1)[1] for u, _ in calls] == ["2027.xlsx", "2026.xlsx"]
    assert all("Mozilla/5.0" in h["User-Agent"] for _, h in calls)


@pytest.mark.parametrize("resp", [_Resp(202, b""), _Resp(403, b"<html>blocked</html>"),
                                  _Resp(200, b"<html>challenge</html>"), _Resp(200, b"")])
def test_hud_fetch_refuses_anything_but_a_200_holding_a_zip(monkeypatch, resp):
    """HUD answers a request it dislikes with HTTP 202 and an empty body; a check on
    the status alone would take that for success (2026-10-07 runner probe)."""
    monkeypatch.setattr(ra.requests, "get", lambda url, headers=None, timeout=None: resp)
    with pytest.raises(RuntimeError, match="not a workbook"):
        ra.fetch_hud_workbook(today=date(2026, 10, 7))


def test_hud_fetch_gives_up_when_no_candidate_year_is_published(monkeypatch):
    monkeypatch.setattr(ra.requests, "get",
                        lambda url, headers=None, timeout=None: _Resp(404, b"<html/>"))
    with pytest.raises(RuntimeError, match="no FMR_2Bed_1983_<year>.xlsx published for 2025-2027"):
        ra.fetch_hud_workbook(today=date(2026, 10, 7))


def test_hud_spec_flows_through_fetch_values_and_rejects_the_percentile_column(monkeypatch):
    spec = SPECS["hud_fmr_honolulu.json"]
    monkeypatch.setattr(ra, "fetch_hud_workbook", lambda: _hud_workbook(HUD_ROWS))
    assert ra.fetch_values(spec, None)[2018] == 2031
    # a column mix-up that put the percentile (40) where the dollars belong
    monkeypatch.setattr(ra, "fetch_hud_workbook", lambda: _hud_workbook({y: 40 for y in HUD_ROWS}))
    with pytest.raises(ValueError, match="outside"):
        ra.fetch_values(spec, None)


# ---------------------------------------------------------------------------
# The bundled files
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("name", sorted(SPECS))
def test_bundled_anchor_is_contiguous_plausible_and_current_enough(name):
    path = ra.ANCHORS_DIR / name
    d = json.loads(path.read_text())
    years = sorted(int(y) for y in d["values_by_year"])
    assert years[0] == ra.FIRST_YEAR
    assert years == list(range(years[0], years[-1] + 1)), "gap in values_by_year"
    ra.check_plausible(name, {int(y): float(v) for y, v in d["values_by_year"].items()})
    assert years[-1] >= 2024
